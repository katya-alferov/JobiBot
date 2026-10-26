"""Extract a structured, source-tracked view of a CV in DOCX format.

The parser only reads. It never modifies the original file.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import docx

DATE_RE = re.compile(r"\(([^()]*\d{4}[^()]*)\)\s*$")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
PHONE_RE = re.compile(r"\(?\+?\d[\d()\s-]{7,}\d")


@dataclass
class Role:
    title: str
    organization: str | None
    dates: str
    start: str | None
    end: str | None
    bullets: list[str] = field(default_factory=list)
    header_text: str = ""

    @property
    def key(self) -> str:
        base = self.organization or self.title
        return re.sub(r"[^a-z0-9]+", "-", base.lower()).strip("-")


@dataclass
class EducationEntry:
    title: str
    institution: str | None
    dates: str
    details: list[str] = field(default_factory=list)
    header_text: str = ""

    @property
    def key(self) -> str:
        return re.sub(r"[^a-z0-9]+", "-", self.title.lower()).strip("-")


@dataclass
class SkillGroup:
    label: str
    items: list[str]
    line: str


@dataclass
class ParsedCV:
    name: str
    contact_line: str
    email: str | None
    phone: str | None
    summary: str
    experience: list[Role]
    education: list[EducationEntry]
    courses: list[str]
    skills: list[SkillGroup]
    languages: list[tuple[str, str]]
    extra_lines: list[str]
    full_text: str

    def all_bullets(self) -> list[str]:
        return [b for r in self.experience for b in r.bullets]


def split_top_level(text: str, sep: str = ",") -> list[str]:
    """Split on separators that are not inside parentheses."""
    parts, depth, buf = [], 0, []
    for ch in text:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        if ch == sep and depth == 0:
            parts.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
    parts.append("".join(buf))
    return [p.strip() for p in parts if p.strip()]


def _is_heading(p) -> bool:
    text = p.text.strip()
    if not text or len(text) > 60:
        return False
    letters = [c for c in text if c.isalpha()]
    bold = any(r.bold for r in p.runs if r.text.strip())
    return bool(letters) and text.upper() == text and bold and not DATE_RE.search(text)


def _split_header(text: str) -> tuple[str, str | None, str, str | None, str | None]:
    m = DATE_RE.search(text)
    dates = m.group(1).strip() if m else ""
    head = text[: m.start()].strip() if m else text.strip()
    if "|" in head:
        title, org = [part.strip() for part in head.split("|", 1)]
    else:
        title, org = head, None
    start = end = None
    if dates:
        parts = [d.strip() for d in re.split(r"\s+-\s+|\s*–\s*", dates)]
        start = parts[0] if parts else None
        end = parts[1] if len(parts) > 1 else None
    return title, org, dates, start, end


def parse_docx(path: Path) -> ParsedCV:
    document = docx.Document(str(path))
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    if not paragraphs:
        raise ValueError(f"{path.name} contains no text")

    name = paragraphs[0].text.strip()
    contact_line = paragraphs[1].text.strip() if len(paragraphs) > 1 else ""
    email_m = EMAIL_RE.search(contact_line)
    phone_m = PHONE_RE.search(contact_line)

    sections: dict[str, list] = {}
    order: list[str] = []
    current = "_preamble"
    for p in paragraphs[2:]:
        if _is_heading(p):
            current = p.text.strip()
            order.append(current)
            sections[current] = []
        else:
            sections.setdefault(current, []).append(p)

    def find_section(*names: str) -> list:
        for heading in order:
            if any(n in heading for n in names):
                return sections[heading]
        return []

    summary = " ".join(p.text.strip() for p in find_section("SUMMARY", "PROFILE"))

    experience: list[Role] = []
    for p in find_section("EXPERIENCE"):
        text = p.text.strip()
        is_bullet = "List" in (p.style.name or "")
        if not is_bullet and DATE_RE.search(text):
            title, org, dates, start, end = _split_header(text)
            experience.append(Role(title, org, dates, start, end, header_text=text))
        elif experience:
            experience[-1].bullets.append(text)

    education: list[EducationEntry] = []
    for p in find_section("EDUCATION"):
        text = p.text.strip()
        if DATE_RE.search(text):
            title, inst, dates, _, _ = _split_header(text)
            education.append(EducationEntry(title, inst, dates, header_text=text))
        elif education:
            education[-1].details.append(text)

    courses = [p.text.strip() for p in find_section("COURSES")]

    skills: list[SkillGroup] = []
    for p in find_section("SKILLS", "KNOWLEDGE"):
        text = p.text.strip()
        if ":" in text:
            label, rest = text.split(":", 1)
            items = [i.rstrip(".") for i in split_top_level(rest) if i.rstrip(".")]
            skills.append(SkillGroup(label.strip(), items, text))
        else:
            skills.append(SkillGroup("Other", [text.rstrip(".")], text))

    languages: list[tuple[str, str]] = []
    extra: list[str] = []
    for p in find_section("LANGUAGES"):
        text = p.text.strip()
        pairs = re.findall(r"([A-Za-z֐-׿]+)\s*:\s*([^|]+)", text)
        if pairs and "Recommend" not in text:
            languages.extend((lang.strip(), level.strip()) for lang, level in pairs)
        else:
            extra.append(text)

    full_text = "\n".join(p.text for p in document.paragraphs)
    return ParsedCV(
        name=name,
        contact_line=contact_line,
        email=email_m.group(0) if email_m else None,
        phone=phone_m.group(0).strip() if phone_m else None,
        summary=summary,
        experience=experience,
        education=education,
        courses=courses,
        skills=skills,
        languages=languages,
        extra_lines=extra,
        full_text=full_text,
    )


def compare_cvs(a: ParsedCV, b: ParsedCV) -> dict:
    """Line-level differences between two CV versions (a = reference, b = candidate)."""
    a_lines = {line.strip() for line in a.full_text.splitlines() if line.strip()}
    b_lines = {line.strip() for line in b.full_text.splitlines() if line.strip()}
    a_skills = {i for g in a.skills for i in g.items}
    b_skills = {i for g in b.skills for i in g.items}
    return {
        "only_in_reference": sorted(a_lines - b_lines),
        "only_in_candidate": sorted(b_lines - a_lines),
        "skills_added": sorted(b_skills - a_skills),
        "skills_removed": sorted(a_skills - b_skills),
        "roles_reference": [(r.title, r.organization, r.dates) for r in a.experience],
        "roles_candidate": [(r.title, r.organization, r.dates) for r in b.experience],
    }
