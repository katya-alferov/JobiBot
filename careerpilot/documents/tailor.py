"""Tailored CV variants built only from verified facts, with a change log and fact validation.

The master CV file is never modified. Each variant is a new, read-only file with its own version ID.
"""

from __future__ import annotations

import re
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx.shared import Pt
from sqlalchemy.orm import Session

from careerpilot.config import Settings
from careerpilot.profile.cv_parser import ParsedCV, Role, _split_header, parse_docx
from careerpilot.profile.service import ProfileView, _alias_pattern, make_read_only, sha256_file, vocabulary
from careerpilot.storage.audit import audit
from careerpilot.storage.models import JobMatch, JobPosting, ResumeChange, ResumeVariant, utcnow

MARKETING_ROLE_HINTS = ("marketing", "content", "email")
DATA_ROLE_HINTS = ("etl", "bi", "developer", "data")
TRACK_FOCUS = {
    "data_bi": "data, ETL and BI",
    "ai_automation": "data, automation and applied AI",
    "grc_security": "data validation, controls and production reliability",
    "marketing": "digital marketing and analytics",
    "other": "data and business systems",
}


class FactValidationError(ValueError):
    pass


@dataclass
class ChangeEntry:
    section: str
    change_type: str
    original: str
    revised: str
    reason: str
    fact_source: str | None = None
    material: bool = False


@dataclass
class CVPlan:
    contact_line: str
    summary: str
    summary_sources: list[str]
    experience: list[Role]
    skills: list[tuple[str, list[str]]]
    education_lines: list[tuple[str, list[str]]]
    courses: list[str]
    languages_line: str
    extra_lines: list[str]
    section_order: list[str]
    changes: list[ChangeEntry] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    track: str = ""


def _relevance(text: str, skills: list[str], vocab: dict) -> int:
    return sum(1 for s in skills if any(_alias_pattern(a).search(text) for a in vocab.get(s, [s])))


def _role_line(role: Role) -> str:
    org = f" at {role.organization.split('(')[0].strip()}" if role.organization else ""
    return f"{role.title}{org} ({role.dates})"


def build_plan(view: ProfileView, job: JobPosting, match: JobMatch | None, track: str,
               settings: Settings) -> CVPlan:
    parsed: ParsedCV = view.parsed_cv
    if parsed is None:
        raise RuntimeError("Master CV file not found; import it first")
    vocab = vocabulary(settings)
    target_skills = list(match.matching_skills) if match else []
    changes: list[ChangeEntry] = []
    warnings: list[str] = []
    src = view.master and f"master_cv:{view.master.sha256[:12]}"

    # Contact line: master contact plus links only when verified.
    contact = parsed.contact_line
    links = []
    p = view.profile
    if p.linkedin_profile_url and p.linkedin_profile_verified:
        links.append(p.linkedin_profile_url.replace("https://", "").replace("www.", ""))
    elif p.linkedin_profile_url:
        warnings.append("LinkedIn URL not yet verified as reachable; left out of the CV")
    if p.personal_website_url and p.personal_website_verified:
        links.append(p.personal_website_url.replace("https://", "").rstrip("/"))
    elif p.personal_website_url:
        warnings.append("Website URL not yet verified as reachable; left out of the CV")
    else:
        warnings.append("Personal website URL not supplied yet")
    if links:
        contact = f"{contact}  |  " + "  |  ".join(links)
        changes.append(ChangeEntry("contact", "add_link", parsed.contact_line, contact,
                                   "Add verified portfolio links", "profile.yaml (verified URL)"))

    # Summary: a factual opening sentence from role headers, then the master summary unchanged.
    roles = parsed.experience
    if track == "marketing":
        lead_roles = [r for r in roles if any(h in r.title.lower() for h in MARKETING_ROLE_HINTS)][:2]
        current = roles[0] if roles else None
        intro = "Digital marketing background as " + " and ".join(_role_line(r) for r in lead_roles)
        if current:
            intro += f", currently {_role_line(current)}"
        intro += "."
        sources = [f"{src}#experience/{r.key}" for r in lead_roles + ([current] if current else [])]
    else:
        lead_roles = [r for r in roles if any(h in r.title.lower() for h in DATA_ROLE_HINTS)][:2]
        intro = " and ".join(_role_line(r) for r in lead_roles)
        intro = (intro[0].upper() + intro[1:] + ".") if intro else ""
        sources = [f"{src}#experience/{r.key}" for r in lead_roles]
    skill_items = _cv_skill_names(parsed, target_skills, vocab)
    if skill_items:
        intro += " Relevant skills: " + ", ".join(skill_items[:6]) + "."
        sources.append(f"{src}#skills")
    summary = f"{intro} {parsed.summary}".strip()
    changes.append(ChangeEntry("summary", "reword", parsed.summary, summary,
                               f"Open with verified roles and tools relevant to {TRACK_FOCUS.get(track, track)}",
                               "; ".join(sources), material=True))

    # Experience: same roles, titles and dates; bullets reordered by relevance.
    experience = []
    for role in roles:
        ordered = sorted(role.bullets, key=lambda b: -_relevance(b, target_skills, vocab))
        if ordered != role.bullets:
            changes.append(ChangeEntry(f"experience/{role.key}", "reorder", "\n".join(role.bullets),
                                       "\n".join(ordered), "Put bullets matching the job's skills first",
                                       f"{src}#experience/{role.key}"))
        experience.append(Role(role.title, role.organization, role.dates, role.start, role.end, ordered,
                               role.header_text))

    # Candidate-confirmed roles that are not on the master CV, only for the tracks they allow.
    for fact in _track_scoped_roles(view, track):
        title, org, dates, start, end = _split_header(fact.value)
        added = Role(title, org, dates, start, end, [], fact.value)
        pos = next((i for i, r in enumerate(experience) if _sort_key(r) <= _sort_key(added)), len(experience))
        experience.insert(pos, added)
        changes.append(ChangeEntry("experience", "add_verified_role", "", fact.value,
                                   f"Candidate-confirmed role, included for {TRACK_FOCUS.get(track, track)} roles only",
                                   fact.source, material=True))

    # Skills: groups and items reordered; nothing added.
    groups = [(g.label, list(g.items)) for g in parsed.skills]
    if track == "marketing":
        omitted = [g for g in groups if g[0].lower().startswith("development environments")]
        for label, items in omitted:
            changes.append(ChangeEntry("skills", "omit", f"{label}: {', '.join(items)}", "",
                                       "Less relevant for a marketing role", f"{src}#skills"))
        groups = [g for g in groups if g not in omitted]
    reordered = sorted(groups, key=lambda g: -_relevance(" ".join(g[1]), target_skills, vocab))
    reordered = [(label, sorted(items, key=lambda i: -_relevance(i, target_skills, vocab)))
                 for label, items in reordered]
    if [g[0] for g in reordered] != [g[0] for g in groups]:
        changes.append(ChangeEntry("skills", "reorder", ", ".join(g[0] for g in groups),
                                   ", ".join(g[0] for g in reordered), "Lead with skill groups the job asks for",
                                   f"{src}#skills"))

    order = ["summary", "experience", "skills", "education", "courses", "languages"]
    if track in {"data_bi", "ai_automation"} and target_skills:
        order = ["summary", "skills", "experience", "education", "courses", "languages"]
        changes.append(ChangeEntry("layout", "reorder", "Summary, Experience, Skills, Education",
                                   "Summary, Skills, Experience, Education",
                                   "Technical reviewers scan skills first", None))

    # Degrees whose completion is unconfirmed are left out until the candidate confirms them.
    education = [(e.header_text, e.details) for e in parsed.education]
    for subject in _unconfirmed_completions(view):
        dropped = [e for e in education if subject in e[0].lower()]
        for header, details in dropped:
            changes.append(ChangeEntry("education", "omit", "\n".join([header, *details]), "",
                                       "Completion not yet confirmed by the candidate",
                                       f"{src}#education", material=True))
            warnings.append(f"Left out until completion is confirmed: {header}")
        education = [e for e in education if e not in dropped]

    return CVPlan(
        contact_line=contact, summary=summary, summary_sources=sources, experience=experience,
        skills=reordered, education_lines=education,
        courses=list(parsed.courses),
        languages_line=" | ".join(f"{lang}: {lvl}" for lang, lvl in parsed.languages),
        extra_lines=list(parsed.extra_lines), section_order=order, changes=changes, warnings=warnings,
        track=track,
    )


def _track_scoped_roles(view: ProfileView, track: str) -> list:
    return [f for f in view.verified if f.category == "employment" and (f.detail or {}).get("tracks")
            and track in f.detail["tracks"]]


def _year(text: str | None) -> int:
    if not text or text.lower().startswith("present"):
        return 9999
    m = re.search(r"(\d{4})", text)
    return int(m.group(1)) if m else 0


def _sort_key(role: Role) -> tuple[int, int]:
    return (_year(role.end or role.start), _year(role.start))


def _unconfirmed_completions(view: ProfileView) -> list[str]:
    return [f.key.removesuffix("-completion").replace("-", " ") for f in view.unverified
            if f.category == "education" and f.key.endswith("-completion")]


def _cv_skill_names(parsed: ParsedCV, target: list[str], vocab: dict) -> list[str]:
    """Skill names exactly as written on the CV that correspond to the job's matched skills."""
    names = []
    for group in parsed.skills:
        for item in group.items:
            if _relevance(item, target, vocab) and item not in names:
                names.append(item)
    return names


def validate_plan(plan: CVPlan, view: ProfileView) -> list[str]:
    """Every factual line must come from a verified fact; unverified facts must not appear."""
    verified_values = {f.value.strip() for f in view.verified}
    verified_skills = {i for f in view.facts_by_category("skill_group") for i in f.detail.get("items", [])}
    errors = []
    for role in plan.experience:
        if role.header_text.strip() not in verified_values:
            errors.append(f"role header not verified: {role.header_text}")
        for b in role.bullets:
            if b.strip() not in verified_values:
                errors.append(f"bullet not verified: {b[:80]}")
    scoped = {f.value.strip(): f.detail["tracks"] for f in view.verified
              if f.category == "employment" and (f.detail or {}).get("tracks")}
    for role in plan.experience:
        allowed = scoped.get(role.header_text.strip())
        if allowed and plan.track not in allowed:
            errors.append(f"role limited to {allowed} used on a {plan.track} CV: {role.header_text}")
    for _, items in plan.skills:
        for item in items:
            if item not in verified_skills:
                errors.append(f"skill not verified: {item}")
    for header, _ in plan.education_lines:
        if header.strip() not in verified_values:
            errors.append(f"education not verified: {header}")
    for subject in _unconfirmed_completions(view):
        if any(subject in h.lower() for h, _ in plan.education_lines):
            errors.append(f"degree with unconfirmed completion presented: {subject}")
    for fact in view.unverified:
        if fact.key.endswith("-completion"):
            continue
        probe = fact.value.split(",")[0].split("(")[0].strip()
        if len(probe) > 12 and probe.lower() in _plan_text(plan).lower():
            errors.append(f"unverified fact used: {fact.key}")
    for hint in re.findall(r"\b\d{2,}%|\$\d|₪\d", plan.summary):
        errors.append(f"summary contains an unsupported figure: {hint}")
    return errors


def _plan_text(plan: CVPlan) -> str:
    parts = [plan.contact_line, plan.summary]
    for r in plan.experience:
        parts += [r.header_text, *r.bullets]
    parts += [f"{label}: {', '.join(items)}" for label, items in plan.skills]
    parts += [h for h, _ in plan.education_lines] + plan.courses + [plan.languages_line]
    return "\n".join(parts)


def render_docx(plan: CVPlan, name: str, path: Path) -> None:
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    def heading(text: str):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(12)

    title = doc.add_paragraph()
    tr = title.add_run(name.upper())
    tr.bold = True
    tr.font.size = Pt(18)
    doc.add_paragraph(plan.contact_line)

    for section in plan.section_order:
        if section == "summary":
            heading("SUMMARY")
            doc.add_paragraph(plan.summary)
        elif section == "experience":
            heading("PROFESSIONAL EXPERIENCE")
            for role in plan.experience:
                p = doc.add_paragraph()
                p.add_run(role.header_text).bold = True
                for b in role.bullets:
                    doc.add_paragraph(b, style="List Bullet")
        elif section == "skills":
            heading("PROFESSIONAL KNOWLEDGE & SKILLS")
            for label, items in plan.skills:
                doc.add_paragraph(f"{label}: {', '.join(items)}", style="List Bullet")
        elif section == "education":
            heading("EDUCATION")
            for header, details in plan.education_lines:
                p = doc.add_paragraph()
                p.add_run(header).bold = True
                for d in details:
                    doc.add_paragraph(d)
        elif section == "courses" and plan.courses:
            heading("RELEVANT COURSES")
            for c in plan.courses:
                doc.add_paragraph(c, style="List Bullet")
        elif section == "languages":
            heading("LANGUAGES & RECOMMENDATIONS")
            doc.add_paragraph(plan.languages_line, style="List Bullet")
            for line in plan.extra_lines:
                doc.add_paragraph(line, style="List Bullet")
    doc.save(str(path))


def convert_to_pdf(docx_path: Path) -> Path | None:
    """Convert with a local LibreOffice install (free). Returns None if unavailable."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent),
                        str(docx_path)], check=True, capture_output=True, timeout=120)
    except (subprocess.SubprocessError, OSError):
        return None
    pdf = docx_path.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def _slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")[:30] or "job"


def create_variant(session: Session, settings: Settings, view: ProfileView, job: JobPosting,
                   match: JobMatch | None, *, make_pdf: bool = True) -> ResumeVariant:
    track = match.track if match else job.career_track
    plan = build_plan(view, job, match, track, settings)
    errors = validate_plan(plan, view)
    if errors:
        raise FactValidationError("; ".join(errors))

    stamp = utcnow().strftime("%Y%m%d%H%M%S%f")[:17]
    version_id = f"cv-{track}-{job.id[-6:]}-{stamp}"
    filename = f"{_slug(view.profile.full_name)}_CV_{track}_{_slug(job.employer)}_{version_id[-17:]}.docx"
    out_dir = settings.data_dir / "cv" / "variants" / version_id
    out_dir.mkdir(parents=True, exist_ok=False)
    docx_path = out_dir / filename
    render_docx(plan, view.profile.full_name, docx_path)

    # Validate the written file opens and contains the intended content.
    reread = parse_docx(docx_path)
    content_checks = {
        "opens": True,
        "name_present": reread.name.lower() == view.profile.full_name.lower(),
        "roles_match_plan": [r.header_text for r in reread.experience]
        == [r.header_text for r in plan.experience],
        "summary_present": reread.summary.strip() == plan.summary.strip(),
        "bullet_count_match": len(reread.all_bullets()) == sum(len(r.bullets) for r in plan.experience),
    }
    pdf_path = convert_to_pdf(docx_path) if make_pdf else None
    if pdf_path:
        make_read_only(pdf_path)
    make_read_only(docx_path)
    valid = all(v for v in content_checks.values())

    variant = ResumeVariant(
        version_id=version_id, master_id=view.master.id, job_id=job.id, track=track, language="en",
        profile_version=view.profile.version, docx_filename=filename, docx_path=str(docx_path),
        docx_sha256=sha256_file(docx_path), pdf_path=str(pdf_path) if pdf_path else None,
        pdf_sha256=sha256_file(pdf_path) if pdf_path else None,
        validation={"fact_errors": errors, "content_checks": content_checks, "warnings": plan.warnings,
                    "pdf": "created" if pdf_path else "LibreOffice not available; DOCX only",
                    "language_note": "English only: a Hebrew CV needs a Hebrew master CV from the candidate"
                    if job.language == "he" else None},
        valid=valid,
    )
    session.add(variant)
    session.flush()
    for c in plan.changes:
        session.add(ResumeChange(variant_id=variant.id, section=c.section, change_type=c.change_type,
                                 original_text=c.original, revised_text=c.revised, reason=c.reason,
                                 fact_source=c.fact_source, material=c.material))
    audit(session, "resume_variant_created", entity_type="resume_variant", entity_id=variant.version_id,
          job_id=job.id, track=track, sha256=variant.docx_sha256, valid=valid)
    return variant
