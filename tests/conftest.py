"""Test fixtures. Everything is synthetic: no real CV, no real employer, no network."""

from __future__ import annotations

import shutil
from pathlib import Path

import docx
import pytest

from careerpilot.config import CONFIG_DIR, load_settings
from careerpilot.storage.db import Database

FIXTURES = Path(__file__).parent / "fixtures"
MOCK_DOMAIN = "mock-ats.careerpilot.test"


def make_synthetic_cv(path: Path, *, extra_skill: str | None = None) -> Path:
    d = docx.Document()

    def bold(text):
        p = d.add_paragraph()
        p.add_run(text).bold = True

    def bullet(text):
        d.add_paragraph(text, style="List Bullet")

    bold("DANA EXAMPLE")
    d.add_paragraph("(050) 000-0000  |  dana@example.test")
    bold("SUMMARY")
    d.add_paragraph("Data professional with strong analytical skills.")
    bold("PROFESSIONAL EXPERIENCE")
    bold("ETL & BI Developer | Example Telecom (Billing Team)  (11/2024 - Present)")
    bullet("Develop and maintain ETL processes using Informatica PowerCenter.")
    bullet("Create technical and functional specification documents.")
    bullet("Develop reports and dashboards using SAP BusinessObjects and Power BI.")
    bullet("Write complex SQL on Oracle and Microsoft SQL Server for data validation.")
    bold("ETL Developer | Example Card  (2022 - 2024)")
    bullet("Provided shift-based support for Informatica and Control-M failures.")
    bullet("Created tables, views and stored procedures in Teradata.")
    bullet("Worked with Jenkins and Jira.")
    bold("Marketing Manager | Example Agency  (2020 - 2021)")
    bullet("Managed content writers, PPC specialists and SEO professionals.")
    bold("Freelancer Digital Marketing Analyst  (2017 - 2021)")
    bullet("Tracked KPIs, designed A/B tests and reported on SEO and PPC using Google Analytics.")
    bold("EDUCATION")
    bold("B.A. Computer Science & Business Administration | Example University  (2020 - 2024)")
    d.add_paragraph("GPA: 87")
    bold("RELEVANT COURSES")
    bullet("Database Systems - SQL (Grade: 93)")
    bold("PROFESSIONAL KNOWLEDGE & SKILLS")
    bullet("Programming Languages: SQL, Python, Java")
    bullet("Development Environments: Eclipse, Visual Studio Code")
    bullet("Databases: Teradata, Oracle, Microsoft SQL Server (MSS)")
    bullet("ETL & BI Tools: Informatica PowerCenter, SAP BusinessObjects (BO), Power BI, Control-M")
    bullet("AI Tools & Technologies: Generative AI, Prompt Engineering, AI Agents")
    bullet("Digital Marketing & Analytics: Google Analytics, Meta Ads, SEO, PPC, Email Marketing")
    bullet("Other Tools & Methodologies: Git, Jira, Jenkins" + (f", {extra_skill}" if extra_skill else ""))
    bold("LANGUAGES & RECOMMENDATIONS")
    bullet("Hebrew: Mother tongue | English: Fluent | Russian: Fluent")
    bullet("Recommendations will be provided upon request.")
    d.save(str(path))
    return path


@pytest.fixture
def config_dir(tmp_path) -> Path:
    cfg = tmp_path / "config"
    shutil.copytree(CONFIG_DIR, cfg)
    (cfg / "profile.yaml").write_text(
        'personal_website_url: ""\n'
        'linkedin_profile_url: "https://www.linkedin.com/in/dana-example"\n'
        "additional_facts:\n"
        "  - category: employment\n"
        "    key: side-brand\n"
        '    value: "Founder & Owner, Example Swim Brand (swimwear), 2017 - 2021"\n'
        "    status: needs_confirmation\n"
        '    source: "test"\n',
        encoding="utf-8",
    )
    sources = cfg / "sources.yaml"
    sources.write_text("greenhouse: []\nlever: []\nuse_synthetic_fixtures: false\n", encoding="utf-8")
    return cfg


@pytest.fixture
def make_settings(tmp_path, config_dir, monkeypatch):
    for var in ("ZERO_COST_MODE", "DRY_RUN", "SUBMISSIONS_KILL_SWITCH", "DATABASE_URL", "DATA_DIR",
                "APPROVED_SUBMISSION_DOMAINS", "LLM_PROVIDER", "MAX_DAILY_APPLICATIONS"):
        monkeypatch.delenv(var, raising=False)

    def factory(**overrides):
        values = {"data_dir": tmp_path / "data", "config_dir": config_dir, "database_url": "",
                  "approved_submission_domains": [MOCK_DOMAIN], "generate_pdf": False}
        values.update(overrides)
        return load_settings(env_file=tmp_path / "missing.env", **values)

    return factory


@pytest.fixture
def settings(make_settings):
    return make_settings()


@pytest.fixture
def db(settings):
    return Database(settings, migrate=False)


@pytest.fixture
def cv_path(tmp_path) -> Path:
    return make_synthetic_cv(tmp_path / "Dana_Example_CV.docx")


@pytest.fixture
def profiled_db(db, settings, cv_path):
    from careerpilot.profile.service import get_profile, import_master_cv

    with db.session() as s:
        import_master_cv(s, settings, cv_path)
        p = get_profile(s)
        p.linkedin_profile_verified = True  # tests never hit the network
    return db


@pytest.fixture
def ingested_db(profiled_db, settings):
    from careerpilot.ingestion.pipeline import run_source
    from careerpilot.matching.engine import match_job
    from careerpilot.profile.service import load_profile_view
    from careerpilot.sources.adapters import FixtureSource
    from careerpilot.storage.models import JobPosting
    from sqlalchemy import select

    with profiled_db.session() as s:
        run_source(s, settings, FixtureSource(FIXTURES / "synthetic_jobs.json"))
        view = load_profile_view(s, settings)
        for job in s.scalars(select(JobPosting)):
            match_job(s, job, view, settings)
    return profiled_db


def job_by_source_id(session, source_job_id: str):
    from careerpilot.storage.models import JobSourceListing
    from sqlalchemy import select

    listing = session.scalars(select(JobSourceListing).where(JobSourceListing.source_job_id == source_job_id)).first()
    return listing.job
